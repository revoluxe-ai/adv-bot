"""Запись и чтение .xlsx с оформлением колонок RO/EDIT."""

from __future__ import annotations

import io
from typing import Any, Literal, cast

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.client_bot.user_interface.services.adv_excel_roundtrip.schema import (
    EXCEL_TEMPLATE_VERSION,
    enum_names_for_field,
    parse_header,
)

LEGEND_TEXT = (
    f"Шаблон v{EXCEL_TEMPLATE_VERSION}. Редактируйте только колонки с тех. ключом (EDIT__...). "
    "BOOL: только TRUE или FALSE. Любая ошибка в файле — отклонение импорта целиком."
)

RO_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
EDIT_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
HEADER_RO_FONT = Font(bold=True, color="666666")
HEADER_EDIT_FONT = Font(bold=True, color="006100")

_NumericDvType = Literal["decimal"]
_NumericDvOperator = Literal["greaterThanOrEqual", "greaterThan"]


def write_adv_excel_workbook(headers: list[str], rows: list[dict[str, Any]]) -> bytes:
    """Строка 1 — легенда; строка 2 — заголовки; с строки 3 — данные."""
    df = pd.DataFrame(rows, columns=cast(Any, headers))
    buf = io.BytesIO()
    with pd.ExcelWriter(cast(Any, buf), engine="openpyxl") as writer:
        sheet = "adv_campaigns"
        cast(Any, df).to_excel(writer, sheet_name=sheet, index=False, startrow=1)
        ws = writer.book[sheet]
        ws.cell(row=1, column=1, value=LEGEND_TEXT)
        ncols = len(headers)
        if ncols:
            ws.merge_cells(
                start_row=1,
                end_row=1,
                start_column=1,
                end_column=min(ncols, 16),
            )
        _style_headers(ws, headers, header_row=2)
        nrows = len(rows)
        _apply_bool_validation(ws, headers, data_start_row=3, data_row_count=nrows)
        _apply_enum_validation(writer.book, ws, headers, data_start_row=3, data_row_count=nrows)
        _apply_numeric_validation(ws, headers, data_start_row=3, data_row_count=nrows)
    buf.seek(0)
    return buf.getvalue()


def _style_headers(ws: Any, headers: list[str], header_row: int) -> None:
    """Подсветка RO/EDIT заголовков."""
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        kind, _field, _suf = parse_header(name)
        if kind == "RO":
            cell.fill = RO_FILL
            cell.font = HEADER_RO_FONT
        elif kind == "EDIT":
            cell.fill = EDIT_FILL
            cell.font = HEADER_EDIT_FONT


def _apply_bool_validation(
    ws: Any, headers: list[str], data_start_row: int, data_row_count: int
) -> None:
    """DataValidation list TRUE/FALSE для BOOL-колонок."""
    if data_row_count <= 0:
        return
    last_row = data_start_row + data_row_count - 1
    for col_idx, name in enumerate(headers, start=1):
        _kind, _field, suffix = parse_header(name)
        if suffix != "BOOL":
            continue
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(
            type="list",
            formula1='"TRUE,FALSE"',
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        cast(Any, dv).add(f"{col_letter}{data_start_row}:{col_letter}{last_row}")


def _apply_enum_validation(
    wb: Any, ws: Any, headers: list[str], data_start_row: int, data_row_count: int
) -> None:
    """DataValidation list по скрытому листу ``_validations`` для ENUM-колонок."""
    if data_row_count <= 0:
        return
    last_row = data_start_row + data_row_count - 1
    enum_cols: list[tuple[int, str]] = []
    for col_idx, name in enumerate(headers, start=1):
        _kind, field, suffix = parse_header(name)
        if suffix == "ENUM":
            enum_cols.append((col_idx, field))
    if not enum_cols:
        return

    list_sheet_name = "_validations"
    if list_sheet_name in wb.sheetnames:
        vws = wb[list_sheet_name]
    else:
        vws = wb.create_sheet(list_sheet_name)
    vws.sheet_state = "hidden"

    list_col_idx = 1
    for data_col_idx, field in enum_cols:
        values = enum_names_for_field(field)
        if not values:
            continue
        for i, value in enumerate(values, start=1):
            vws.cell(row=i, column=list_col_idx, value=value)
        enum_col_letter = get_column_letter(list_col_idx)
        formula = f"={list_sheet_name}!${enum_col_letter}$1:${enum_col_letter}${len(values)}"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        col_letter = get_column_letter(data_col_idx)
        cast(Any, dv).add(f"{col_letter}{data_start_row}:{col_letter}{last_row}")
        list_col_idx += 1


def _apply_numeric_validation(
    ws: Any, headers: list[str], data_start_row: int, data_row_count: int
) -> None:
    """Numeric DataValidation для известных EDIT decimal-полей."""
    if data_row_count <= 0:
        return
    last_row = data_start_row + data_row_count - 1
    numeric_rules: dict[str, tuple[_NumericDvType, _NumericDvOperator, str | None]] = {
        "plan_for_sales_for_day": ("decimal", "greaterThanOrEqual", "1"),
        "target_price_for_order": ("decimal", "greaterThanOrEqual", "1"),
        "max_price_for_order": ("decimal", "greaterThanOrEqual", "1"),
        "auto_shelf_max_bid_rub": ("decimal", "greaterThanOrEqual", "1"),
        "auto_shelf_min_bid_distance_rub": ("decimal", "greaterThanOrEqual", "1"),
        "auto_shelf_target_test_impressions": ("decimal", "greaterThanOrEqual", "1"),
        "auto_shelf_target_test_observation_minutes": ("decimal", "greaterThanOrEqual", "60"),
        "auto_shelf_test_share": ("decimal", "greaterThan", "0"),
    }
    for col_idx, name in enumerate(headers, start=1):
        _kind, field, _suffix = parse_header(name)
        if field not in numeric_rules:
            continue
        dv_type, operator, formula1 = numeric_rules[field]
        dv = DataValidation(
            type=dv_type,
            operator=operator,
            formula1=formula1,
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        col_letter = get_column_letter(col_idx)
        cast(Any, dv).add(f"{col_letter}{data_start_row}:{col_letter}{last_row}")


def read_adv_excel_upload(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Читает файл, сохранённый write_adv_excel_workbook (header=1)."""
    bio = io.BytesIO(content)
    df = cast(Any, pd).read_excel(bio, header=1, engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]
    raw_records = cast(list[dict[Any, Any]], df.to_dict(orient="records"))
    records: list[dict[str, Any]] = [
        {str(k): v for k, v in row.items()} for row in raw_records
    ]
    return list(df.columns), records
