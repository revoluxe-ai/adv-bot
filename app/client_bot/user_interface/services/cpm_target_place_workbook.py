"""Export and import workbook helpers for campaign keyword target places."""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from typing import Any, cast

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.client_bot import const
from app.client_bot.user_interface.handlers.adv_setting.adv_setting import (
    ADV_ACCESS_DENIED_MESSAGE,
    CABINET_ACCESS_DENIED_MESSAGE,
    load_adv_for_owner_cabinet_in_session,
    load_wb_user_cabinet_for_adv_in_session,
)
from app.modules.common_slaves.wildberries.models import KeywordSetting
from app.utils.db.session_maker import get_async_session_maker

_sm = get_async_session_maker

_SHEET_NAME = "target_places"
_LEGEND_TEXT = (
    "Шаблон целевых позиций ключей. Меняйте только колонку target_place. "
    "Пустое значение оставляет ключ без изменений. Допустимы только целые числа больше 0."
)
_HEADER_FILL = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_HEADERS = ["keyword_id", "keyword", "current_target_place", "target_place"]


class TargetPlaceAccessDeniedError(Exception):
    """Отказ в записи target_place: кампания недоступна пользователю."""

    def __init__(self, denial_message: str) -> None:
        self.denial_message = denial_message
        super().__init__(denial_message)


@dataclass(slots=True)
class TargetPlaceImportSummary:
    """Result of campaign target place workbook import."""

    processed_rows: int = 0
    updated_rows: int = 0
    skipped_empty_rows: int = 0
    not_found_rows: int = 0
    validation_error_rows: int = 0
    errors: list[str] = field(default_factory=list[str])


def _is_empty(value: Any) -> bool:
    """Return True for blank cell values used by pandas/openpyxl."""
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _parse_positive_int(value: Any, *, field_name: str) -> int:
    """Parse a positive integer from workbook cell value."""
    if isinstance(value, bool):
        raise ValueError(f"{field_name} должно быть целым числом больше 0")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float):
        if math.isnan(value) or not float(value).is_integer():
            raise ValueError(f"{field_name} должно быть целым числом больше 0")
        parsed = int(value)
    elif isinstance(value, str):
        stripped = value.strip()
        if not stripped.isdigit():
            raise ValueError(f"{field_name} должно быть целым числом больше 0")
        parsed = int(stripped)
    else:
        raise ValueError(f"{field_name} должно быть целым числом больше 0")
    if parsed <= 0:
        raise ValueError(f"{field_name} должно быть целым числом больше 0")
    return parsed


async def load_keywords_for_target_place_workbook_in_session(
    session: AsyncSession, adv_original_id: int
) -> list[KeywordSetting]:
    """Load non-blacklist keywords for workbook export (PURE_READ).

    Args:
        session: injectable async session (caller: ``sm(read_only=True)``).
        adv_original_id: campaign AdvSetting primary key.

    Returns:
        KeywordSetting rows ordered by text.
    """
    result = await session.execute(
        select(KeywordSetting)
        .where(KeywordSetting.adv_setting_id == adv_original_id)
        .where(KeywordSetting.hard_status != const.BLACK_LIST)
        .order_by(KeywordSetting.text.asc())
    )
    return list(result.scalars().all())


async def _verify_adv_owner_in_session(
    session: AsyncSession,
    telegram_user_id: int,
    adv_original_id: int,
) -> None:
    """Проверить, что РК принадлежит Telegram-пользователю, иначе raise."""
    wb_user_cabinet = await load_wb_user_cabinet_for_adv_in_session(
        session, telegram_user_id, adv_original_id
    )
    if wb_user_cabinet is None:
        raise TargetPlaceAccessDeniedError(CABINET_ACCESS_DENIED_MESSAGE)
    adv = await load_adv_for_owner_cabinet_in_session(
        session,
        adv_original_id=adv_original_id,
        wb_user_cabinet=wb_user_cabinet,
    )
    if adv is None:
        raise TargetPlaceAccessDeniedError(ADV_ACCESS_DENIED_MESSAGE)


async def apply_target_place_rows_in_session(
    session: AsyncSession,
    adv_original_id: int,
    rows: list[dict[str, Any]],
    summary: TargetPlaceImportSummary,
    *,
    telegram_user_id: int,
) -> TargetPlaceImportSummary:
    """Apply workbook rows to keywords and commit (WRITE).

    Args:
        session: injectable async session (``sm(read_only=False)``).
        adv_original_id: campaign AdvSetting primary key.
        rows: parsed workbook records.
        summary: mutable import summary accumulator.
        telegram_user_id: Telegram id владельца кабинета с РК.

    Returns:
        The same summary instance after mutation/commit.

    Raises:
        TargetPlaceAccessDeniedError: если РК не принадлежит пользователю.
    """
    await _verify_adv_owner_in_session(session, telegram_user_id, adv_original_id)
    result = await session.execute(
        select(KeywordSetting).where(KeywordSetting.adv_setting_id == adv_original_id)
    )
    keywords = result.scalars().all()
    keywords_by_id = {keyword.id: keyword for keyword in keywords}
    keywords_by_text = {keyword.text: keyword for keyword in keywords}

    for row_number, row in enumerate(rows, start=3):
        raw_target_place = row.get("target_place")
        if _is_empty(raw_target_place):
            summary.skipped_empty_rows += 1
            continue
        try:
            target_place = _parse_positive_int(
                raw_target_place, field_name="target_place"
            )
        except ValueError as exc:
            summary.validation_error_rows += 1
            summary.errors.append(f"Строка {row_number}: {exc}")
            continue

        keyword = None
        raw_keyword_id = row.get("keyword_id")
        if not _is_empty(raw_keyword_id):
            try:
                keyword_id = _parse_positive_int(
                    raw_keyword_id, field_name="keyword_id"
                )
            except ValueError as exc:
                summary.validation_error_rows += 1
                summary.errors.append(f"Строка {row_number}: {exc}")
                continue
            keyword = keywords_by_id.get(keyword_id)

        if keyword is None:
            keyword_text = str(row.get("keyword") or "").strip()
            if not keyword_text:
                summary.validation_error_rows += 1
                summary.errors.append(
                    f"Строка {row_number}: не указан keyword для поиска ключа"
                )
                continue
            keyword = keywords_by_text.get(keyword_text)

        if keyword is None:
            summary.not_found_rows += 1
            summary.errors.append(
                f"Строка {row_number}: ключ не найден в текущей РК"
            )
            continue

        keyword.target_place = target_place
        summary.updated_rows += 1

    await session.commit()
    return summary


async def build_target_place_workbook(adv_original_id: int) -> bytes:
    """Build .xlsx template with current target places for one campaign."""
    async with _sm(read_only=True) as session:
        keywords = await load_keywords_for_target_place_workbook_in_session(
            session, adv_original_id
        )

    rows = [
        {
            "keyword_id": keyword.id,
            "keyword": keyword.text,
            "current_target_place": keyword.target_place,
            "target_place": None,
        }
        for keyword in keywords
    ]
    return write_target_place_workbook_rows(rows)


def write_target_place_workbook_rows(rows: list[dict[str, Any]]) -> bytes:
    """Write workbook rows in the target place template format."""
    df = pd.DataFrame(rows, columns=cast(Any, _HEADERS))
    buf = io.BytesIO()
    with pd.ExcelWriter(cast(Any, buf), engine="openpyxl") as writer:
        cast(Any, df).to_excel(writer, sheet_name=_SHEET_NAME, index=False, startrow=1)
        ws: Any = writer.book[_SHEET_NAME]
        ws.cell(row=1, column=1, value=_LEGEND_TEXT)
        ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=len(_HEADERS))
        for col_idx, header in enumerate(_HEADERS, start=1):
            cell = ws.cell(row=2, column=col_idx)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            ws.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 4, 18)
        if rows:
            last_row = len(rows) + 2
            target_col = get_column_letter(_HEADERS.index("target_place") + 1)
            validation = DataValidation(
                type="whole",
                operator="greaterThan",
                formula1="0",
                allow_blank=True,
            )
            ws.add_data_validation(validation)
            cast(Any, validation).add(f"{target_col}3:{target_col}{last_row}")
    buf.seek(0)
    return buf.getvalue()


def read_target_place_workbook(content: bytes) -> list[dict[str, Any]]:
    """Read workbook created by build_target_place_workbook."""
    bio = io.BytesIO(content)
    df = cast(Any, pd).read_excel(bio, header=1, engine="openpyxl")
    df.columns = [str(column).strip() for column in df.columns]
    missing = [header for header in _HEADERS if header not in df.columns]
    if missing:
        raise ValueError(f"В файле нет обязательных колонок: {', '.join(missing)}")
    records = cast(list[dict[Any, Any]], df.to_dict(orient="records"))
    return [{str(key): value for key, value in record.items()} for record in records]


async def apply_target_place_workbook(
    adv_original_id: int, content: bytes, *, telegram_user_id: int
) -> TargetPlaceImportSummary:
    """Apply target_place values from workbook to keywords of one campaign.

    Args:
        adv_original_id: campaign AdvSetting primary key.
        content: xlsx bytes from Telegram upload.
        telegram_user_id: Telegram id владельца кабинета с РК.

    Raises:
        TargetPlaceAccessDeniedError: если РК не принадлежит пользователю.
    """
    rows = read_target_place_workbook(content)
    summary = TargetPlaceImportSummary(processed_rows=len(rows))
    async with _sm(read_only=False) as session:
        summary = await apply_target_place_rows_in_session(
            session,
            adv_original_id,
            rows,
            summary,
            telegram_user_id=telegram_user_id,
        )
    return summary
